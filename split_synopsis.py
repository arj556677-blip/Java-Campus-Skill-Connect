"""
Split Campus_Skill_Connect_Synopsis.docx into 3 documents:
  - Phase1_Synopsis.docx  (Phase 1 – Core Java and OOP)
  - Phase2_Synopsis.docx  (Phase 2 – Java Structure)
  - Phase3_Synopsis.docx  (Phase 3 – Database)

Run this from your PowerShell/CMD terminal in the project folder:
    python split_synopsis.py

Requirements: pip install python-docx
"""

import copy, sys, subprocess, os

# Auto-install python-docx if missing
try:
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from lxml import etree
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "lxml"])
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    from lxml import etree

# ── Paths ────────────────────────────────────────────────────────────────────
BASE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(BASE, "Campus_Skill_Connect_Synopsis.docx")

# ── Phase content for Section 5 replacement ──────────────────────────────────
PHASES = {
    1: {
        "outfile": os.path.join(BASE, "Phase1_Synopsis.docx"),
        "content": [
            ("bold_para", "Phase 1 – Core Java and OOP"),
            ("para",
             "In this phase, the core model classes of the application were designed and implemented "
             "using Java Object-Oriented Programming principles. The following classes were created:"),
            ("bullet",
             "User – Represents a registered student with fields for name, email, password hash, course, "
             "year of study, and bio. Implements encapsulation through private fields with public getters "
             "and setters, a parameterized constructor for registration, and a derived method getFullName() "
             "that computes the full name from first and last name fields."),
            ("bullet",
             "Skill – Represents a skill listed by a student, with fields for skill name, proficiency level "
             "(using a ProficiencyLevel enum: BEGINNER, INTERMEDIATE, ADVANCED), and an optional description. "
             "Includes a parameterized constructor and convenience setter to parse the proficiency level from a String."),
            ("bullet",
             "SkillRequest – Represents a collaboration or skill-exchange request between two students. Contains "
             "fields for sender (fromUserId), receiver (toUserId), the target skill (skillId), an optional message, "
             "and a status field (using a RequestStatus enum: PENDING, ACCEPTED, REJECTED). Includes a parameterized "
             "constructor and a convenience setter to parse the status from a String."),
            ("bullet",
             "Institution – Represents a participating college or institution with fields for name, city, and state."),
            ("bullet",
             "Campus – Represents a campus of an institution, with fields for name, city, and a foreign key "
             "reference to the parent institution."),
            ("para",
             "All classes were organized with no-argument constructors (required for DAO mapping from database "
             "ResultSets), parameterized constructors for creating new objects in code, and full getter/setter "
             "method sets. A toString() method was included in each class for readable debugging output."),
            ("para",
             "The ProficiencyLevel and RequestStatus enumerations were implemented as separate enum types with "
             "a fromString() factory method to safely parse string values from database columns and form inputs, "
             "preventing IllegalArgumentException errors from invalid strings."),
        ]
    },
    2: {
        "outfile": os.path.join(BASE, "Phase2_Synopsis.docx"),
        "content": [
            ("bold_para", "Phase 2 – Java Structure"),
            ("para",
             "In this phase, the codebase was organized into a clean, maintainable package structure and Java "
             "structural concepts were applied to the project architecture."),
            ("bold_para", "Package Organization"),
            ("para", "All classes were distributed across the following packages:"),
            ("bullet",
             "com.campusskillconnect.model – Domain entity classes (User, Skill, SkillRequest, Institution, "
             "Campus, and the ProficiencyLevel and RequestStatus enums)."),
            ("bullet",
             "com.campusskillconnect.dao – Data Access Object classes responsible for all database interactions "
             "(UserDAO, SkillDAO, SkillRequestDAO, InstitutionDAO, DatabaseConnection)."),
            ("bullet",
             "com.campusskillconnect.servlet – Java Servlet classes that handle HTTP requests and responses."),
            ("bullet",
             "com.campusskillconnect.filter – Servlet filters for cross-cutting concerns (AuthFilter for session security)."),
            ("bullet",
             "com.campusskillconnect.util – Utility classes shared across the application (PasswordUtil for BCrypt operations)."),
            ("bullet",
             "com.campusskillconnect.exception – Custom exception classes (AppException)."),
            ("bold_para", "Interface – Connectable"),
            ("para",
             "A Connectable interface was introduced in the DAO layer to define a standard contract for any class "
             "that provides a database connection. The DatabaseConnection class implements this interface, providing "
             "a getConnection() method that returns a java.sql.Connection. This demonstrates the Java interface "
             "concept: any future connection provider can be substituted without changing the DAO classes."),
            ("bold_para", "Custom Exception Handling – AppException"),
            ("para",
             "A custom checked exception class AppException was created, extending Exception. All DAO methods "
             "declare throws AppException in their signatures, so the calling servlet always handles database and "
             "business logic errors explicitly. AppException wraps lower-level SQLExceptions and provides a clean "
             "error message to the presentation layer. Two constructors were provided: one accepting only a message "
             "string, and one accepting both a message and a cause (the wrapped exception), following standard Java "
             "exception chaining practice."),
        ]
    },
    3: {
        "outfile": os.path.join(BASE, "Phase3_Synopsis.docx"),
        "content": [
            ("bold_para", "Phase 3 – Database"),
            ("para",
             "In this phase, the MySQL relational database was designed, created, and connected to the Java "
             "application using JDBC. A full Data Access Object (DAO) layer was implemented to handle all SQL "
             "operations in one place."),
            ("bold_para", "Database Schema"),
            ("para", "The following tables were created in the campusskillconnect MySQL database:"),
            ("bullet",
             "institutions – Stores college/institution records (id, name, city, state, created_at)."),
            ("bullet",
             "campuses – Stores campus records linked to an institution via institution_id foreign key "
             "(id, institution_id, name, city, created_at)."),
            ("bullet",
             "users – Stores registered student accounts (id, email, password_hash, first_name, last_name, "
             "institution_id, campus_id, course, year_of_study, bio, created_at, updated_at). The email column "
             "has a UNIQUE constraint to prevent duplicate registrations."),
            ("bullet",
             "skills – Stores skills listed by students (id, user_id, skill_name, proficiency_level ENUM, "
             "description, created_at). The user_id column is a foreign key referencing users.id."),
            ("bullet",
             "skill_requests – Stores collaboration requests between students (id, from_user_id, to_user_id, "
             "skill_id, message, status ENUM, created_at, updated_at). Both from_user_id and to_user_id are "
             "foreign keys referencing users.id, and skill_id references skills.id."),
            ("bold_para", "JDBC Connection – DatabaseConnection"),
            ("para",
             "The DatabaseConnection class implements the Connectable interface and loads database connection "
             "parameters (URL, username, password) from a db.properties file stored in src/main/resources. The "
             "properties file is read using Class.getResourceAsStream() so that credentials are not hardcoded in "
             "the Java source. The getConnection() method uses DriverManager.getConnection() to obtain a live "
             "java.sql.Connection, which is passed to each DAO method via try-with-resources to ensure automatic closing."),
            ("bold_para", "Data Access Objects (DAOs)"),
            ("para",
             "Each database table has a corresponding DAO class. All SQL queries use PreparedStatement with '?' "
             "placeholders throughout, preventing SQL injection. No user-supplied input is ever concatenated "
             "directly into a SQL string."),
            ("bullet",
             "UserDAO – Provides insertUser() for registration (with duplicate-email detection), findByEmail() "
             "for login credential lookup, findById() for loading a full profile with institution and campus names "
             "via JOIN, updateProfile() for saving profile edits, and searchBySkill() for finding students by "
             "skill keyword using a LIKE query with wildcard matching."),
            ("bullet",
             "SkillDAO – Provides addSkill() with RETURN_GENERATED_KEYS to retrieve the auto-assigned ID, "
             "getSkillsByUserId() to list all skills for a student, findById() for loading a single skill, "
             "updateSkill() to edit skill details, and deleteSkill() to remove a skill."),
            ("bullet",
             "SkillRequestDAO – Provides sendRequest() to insert a new collaboration request, "
             "getReceivedRequests() to load a student's inbox with sender name and skill name via JOIN, "
             "getSentRequests() to load outgoing requests, updateStatus() to accept or reject a request, "
             "and findById() for loading a single request."),
            ("bullet",
             "InstitutionDAO – Provides getAllInstitutions() to populate the registration form's institution "
             "dropdown, and getCampusesByInstitutionId() to load the campuses for a selected institution "
             "(used by the AJAX campus loader on the registration page)."),
            ("bold_para", "Password Security – jBCrypt"),
            ("para",
             "Student passwords are never stored in plain text. During registration, the "
             "PasswordUtil.hashPassword() method applies the BCrypt hashing algorithm (via the jBCrypt library) "
             "to the submitted password before it is inserted into the users table. During login, "
             "PasswordUtil.verifyPassword() uses BCrypt's checkpw() method to compare the submitted password "
             "against the stored hash without ever decrypting it. This is the industry-standard approach for "
             "password storage."),
        ]
    }
}


# ── Helpers ───────────────────────────────────────────────────────────────────

def para_text(p):
    return "".join(r.text for r in p.runs)


def clone_p(src_para, dst_doc):
    """Clone a paragraph element into dst_doc's body."""
    new_p = copy.deepcopy(src_para._p)
    dst_doc.element.body.append(new_p)


def add_para(doc, text, bold=False, bullet=False):
    """Add a paragraph with optional bold formatting or bullet style."""
    if bullet:
        try:
            p = doc.add_paragraph(style="List Bullet")
        except Exception:
            p = doc.add_paragraph()
            p._p.get_or_add_pPr()
    else:
        p = doc.add_paragraph()

    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


# ── Main builder ──────────────────────────────────────────────────────────────

def build_doc(src_doc, phase_num, phase_data):
    paragraphs = src_doc.paragraphs

    # Locate Section 5 and Section 6 paragraph indices
    sec5_idx = next((i for i, p in enumerate(paragraphs)
                     if "5." in p.text and "Methodology" in p.text), None)
    sec6_idx = next((i for i, p in enumerate(paragraphs)
                     if "6." in p.text and "Expected" in p.text), None)

    if sec5_idx is None:
        raise ValueError("Cannot find Section 5 heading in source document")

    dst = Document()

    # Remove the default empty paragraph Word adds
    for p in dst.paragraphs:
        p._element.getparent().remove(p._element)

    # 1. Copy everything BEFORE Section 5
    for i in range(sec5_idx):
        clone_p(paragraphs[i], dst)

    # 2. Clone Section 5 heading paragraph (preserves style/formatting)
    clone_p(paragraphs[sec5_idx], dst)

    # 3. Add phase-specific content
    for (kind, text) in phase_data["content"]:
        if kind == "bullet":
            add_para(dst, text, bold=False, bullet=True)
        elif kind == "bold_para":
            add_para(dst, text, bold=True, bullet=False)
        else:  # "para"
            add_para(dst, text, bold=False, bullet=False)

    # 4. Add a blank spacer paragraph
    dst.add_paragraph()

    # 5. Copy from Section 6 onwards (Expected Outcomes, Scope Note)
    if sec6_idx is not None:
        for i in range(sec6_idx, len(paragraphs)):
            clone_p(paragraphs[i], dst)

    return dst


def main():
    print(f"Reading: {SRC}")
    if not os.path.exists(SRC):
        print(f"ERROR: Source file not found: {SRC}")
        sys.exit(1)

    src_doc = Document(SRC)

    print(f"Paragraphs found: {len(src_doc.paragraphs)}")
    for i, p in enumerate(src_doc.paragraphs):
        t = p.text.strip()
        if t:
            print(f"  [{i:3d}] {t[:90]}")

    for phase_num, phase_data in PHASES.items():
        print(f"\nBuilding Phase {phase_num} document...")
        dst = build_doc(src_doc, phase_num, phase_data)
        dst.save(phase_data["outfile"])
        print(f"  Saved → {phase_data['outfile']}")

    print("\nAll 3 documents created successfully.")


if __name__ == "__main__":
    main()
